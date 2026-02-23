"""트래블월렛 엑셀 파일 파서."""

import io
import re
from datetime import datetime
from pathlib import Path

import msoffcrypto
import pandas as pd

from . import config


def _parse_amount(raw: str | int | float | None) -> float:
    """'+ 10,000' or '- 44,000' 같은 금액 문자열을 float으로 변환."""
    if raw is None:
        return 0.0
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).replace(",", "").replace(" ", "")
    # "+10000" → 10000, "-44000" → -44000
    return float(s)


def _parse_krw(raw: str | int | float | None) -> float:
    """원화 금액 파싱. 콤마 제거 후 float 변환."""
    if raw is None:
        return 0.0
    if isinstance(raw, (int, float)):
        return float(raw)
    return float(str(raw).replace(",", "").replace(" ", ""))


def decrypt_workbook(filepath: str | Path, password: str) -> io.BytesIO:
    """암호화된 엑셀 파일을 복호화하여 BytesIO로 반환."""
    with open(filepath, "rb") as f:
        office_file = msoffcrypto.OfficeFile(f)
        office_file.load_key(password=password)
        decrypted = io.BytesIO()
        office_file.decrypt(decrypted)
        decrypted.seek(0)
        return decrypted


def parse(
    filepath: str | Path,
    password: str | None = None,
    start_date: datetime | None = None,
    end_date: datetime | None = None,
) -> pd.DataFrame:
    """트래블월렛 엑셀 파일을 파싱하여 통합 DataFrame을 반환.

    Returns:
        DataFrame with columns: date, time, source, type, merchant,
                                currency, amount, krw_amount, category
    """
    filepath = Path(filepath)
    if not filepath.exists():
        raise FileNotFoundError(f"파일을 찾을 수 없습니다: {filepath}")

    # 암호화된 파일 복호화
    if password:
        buf = decrypt_workbook(filepath, password)
    else:
        buf = filepath

    import openpyxl
    wb = openpyxl.load_workbook(buf, data_only=True)
    ws = wb[wb.sheetnames[0]]

    # 컬럼 레터 매핑: 1→A, 2→B, ...
    col_letters = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

    rows = []
    for row in ws.iter_rows(min_row=config.TW_DATA_START_ROW, max_row=ws.max_row):
        cells = {}
        for cell in row:
            idx = cell.column - 1
            if idx < len(col_letters):
                cells[col_letters[idx]] = cell.value

        raw_date = cells.get("B")
        if not raw_date or not isinstance(raw_date, str):
            continue
        # "2026.02.20" → date
        match = re.match(r"(\d{4})\.(\d{2})\.(\d{2})", raw_date)
        if not match:
            continue

        dt = datetime(int(match.group(1)), int(match.group(2)), int(match.group(3)))

        # 시간 정보가 있으면 결합
        raw_time = str(cells.get("C", "") or "").strip()
        if raw_time:
            try:
                t = datetime.strptime(raw_time, "%H:%M:%S")
                dt = dt.replace(hour=t.hour, minute=t.minute, second=t.second)
            except ValueError:
                pass

        # 기간 필터링 (None이면 해당 방향 제한 없음)
        if start_date and dt < start_date:
            continue
        if end_date and dt > end_date:
            continue

        raw_type = cells.get("D", "")
        raw_amount = _parse_amount(cells.get("H"))
        raw_krw = _parse_krw(cells.get("M"))

        # 지출은 양수로 통일 (원본에서 -는 지출)
        amount = abs(raw_amount)
        krw_amount = abs(raw_krw)

        # 충전/환불은 지출이 아니므로 부호로 구분
        is_expense = any(t in str(raw_type) for t in ["결제", "payment", "ATM", "atm", "N빵", "split"])

        rows.append({
            "date": dt,
            "time": cells.get("C", ""),
            "source": "트래블월렛",
            "type": str(raw_type),
            "merchant": str(cells.get("J", "")),
            "currency": str(cells.get("G", "")),
            "amount": amount if is_expense else -amount,
            "krw_amount": krw_amount if is_expense else -krw_amount,
            "category": "",
        })

    wb.close()

    df = pd.DataFrame(rows, columns=config.UNIFIED_COLUMNS)
    df["date"] = pd.to_datetime(df["date"])
    df.sort_values(["date", "time"], inplace=True)
    df.reset_index(drop=True, inplace=True)
    return df
